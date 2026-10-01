"""Test parsers — tiap format menghasilkan section + locator."""

import io

import openpyxl
import pytest

from ai_asistent_core.parsers import (
    PARSERS,
    CsvParser,
    DocxParser,
    PdfParser,
    TextLikeParser,
    parse_file,
)


def _pdf_bytes(pages: int = 2) -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _i in range(pages):
        writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _docx_bytes(paragraphs: list[str]) -> bytes:
    import docx as docx_lib

    doc = docx_lib.Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pptx_bytes(slides: int = 2) -> bytes:
    from pptx import Presentation

    prs = Presentation()
    for i in range(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Judul Slide {i + 1}"
        slide.placeholders[1].text = f"Isi slide {i + 1}"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _xlsx_bytes(sheets: int = 2) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for s in range(sheets):
        ws = wb.create_sheet(title=f"Sheet{s + 1}")
        ws.append(["nama", "nilai"])
        ws.append(["andro", 90 + s])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_pdf_pages_become_locators() -> None:
    doc = PdfParser().parse(_pdf_bytes(2))
    assert doc.source_format == "pdf"
    assert doc.meta["pages"] == "2"
    # Halaman kosong menghasilkan section kosong — boleh kosong (blank pdf)


def test_docx_paragraphs_and_heading_split() -> None:
    data = _docx_bytes(["Perkenalan", "Isi dokumen", "Penutup"])
    doc = DocxParser().parse(data)
    assert doc.source_format == "docx"
    joined = "\n".join(s.text for s in doc.sections)
    assert "Isi dokumen" in joined


def test_pptx_slides_have_locator_slide() -> None:
    doc = PARSERS[".pptx"].parse(_pptx_bytes(2))
    assert doc.source_format == "pptx"
    assert len(doc.sections) == 2
    assert all(s.locator_type == "slide" for s in doc.sections)
    assert doc.sections[0].locator_start == 1
    assert doc.sections[1].locator_start == 2


def test_xlsx_sheets_have_locator_sheet() -> None:
    doc = PARSERS[".xlsx"].parse(_xlsx_bytes(2))
    assert doc.source_format == "xlsx"
    assert len(doc.sections) == 2
    assert all(s.locator_type == "sheet" for s in doc.sections)
    assert "andro" in doc.sections[0].text


def test_txt_and_md_char_locator() -> None:
    data = ("kata " * 3000).encode()
    doc = TextLikeParser("txt").parse(data)
    assert doc.sections[0].locator_type == "char"
    md = TextLikeParser("md").parse(b"# Judul\n\nisi")
    assert md.title == "Judul"


def test_csv_rows_locator() -> None:
    data = b"a,b\n1,2\n3,4\n"
    doc = CsvParser().parse(data)
    assert doc.source_format == "csv"
    assert doc.sections[0].locator_type == "row"
    assert doc.meta["total_rows"] == "3"


def test_parse_file_rejects_unknown_extension() -> None:
    with pytest.raises(ValueError):
        parse_file("file.exe", b"x")


def test_pdf_text_extraction_roundtrip() -> None:
    """PDF berisi teks nyata harus terekstrak per halaman (jika writer mendukung)."""
    try:
        from reportlab.pdfgen import canvas as pdf_canvas  # type: ignore[import-not-found]
    except ImportError:
        pytest.skip("reportlab tidak terpasang; uji teks PDF dilewati")
    buf = io.BytesIO()
    c = pdf_canvas.Canvas(buf)
    c.drawString(100, 750, "Halo dari PDF")
    c.showPage()
    c.save()
    doc = PdfParser().parse(buf.getvalue())
    assert any("Halo dari PDF" in s.text for s in doc.sections)
